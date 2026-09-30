#!/usr/bin/env python3
"""Build the reviewed interim Hex snapshots from official downloads (requires openpyxl).

Usage: python3 scripts/uk_aq_build_population_snapshot.py SOURCE_DIRECTORY
See data/population/README.md for source URLs and selection details. No API/database writes.
"""

import argparse
import csv
import datetime
import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def row(code, value, year):
    number = int(value)
    if number <= 0 or float(value) != number:
        raise ValueError(f"Invalid population for {code}: {value}")
    return dict(geo_code=code, population_value=number,
                reference_date=datetime.date(int(year), 6, 30).isoformat())


def nomis(source, codes, year, geography_type=None):
    result = []
    with source.open(encoding="utf-8-sig", newline="") as handle:
        for item in csv.DictReader(handle):
            if item["GEOGRAPHY_CODE"] not in codes:
                continue
            if (item["GENDER_NAME"] != "Total" or item["C_AGE_NAME"] != "All Ages"
                    or item["MEASURES"] != "20100" or item["OBS_STATUS"] != "A"
                    or item["DATE"] != str(year)
                    or (geography_type and item["GEOGRAPHY_TYPECODE"] != geography_type)):
                raise ValueError(f"Unexpected Nomis dimensions: {item}")
            result.append(row(item["GEOGRAPHY_CODE"], item["OBS_VALUE"], item["DATE"]))
    return result


def nisra(source, codes, geography, expected_year):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        rows = workbook["Flat"].values
        if next(rows) != ("area", "area_code", "area_name", "year", "type", "MYE"):
            raise ValueError("NISRA Flat columns changed")
        latest = {}
        seen = set()
        for area, code, name, year, rounding, value in rows:
            if area != geography or rounding != "Unrounded" or code not in codes:
                continue
            if (code, year) in seen:
                raise ValueError(f"Duplicate NISRA row: {code}, {year}")
            seen.add((code, year))
            candidate = row(code, value, year)
            if code not in latest or candidate["reference_date"] > latest[code]["reference_date"]:
                latest[code] = candidate
        if any(r["reference_date"] != f"{expected_year}-06-30" for r in latest.values()):
            raise ValueError("NISRA reference period changed; review the source")
        return list(latest.values())
    finally:
        workbook.close()


def nrs(source, codes):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["Table 1"]
        if not sheet["A1"].value.endswith("mid-2025"):
            raise ValueError("NRS reference period changed; review the source")
        rows = list(sheet.iter_rows(min_row=4, max_col=5, values_only=True))
        if rows[0] != ("Area name", "Area code", "Area type", "Sex", "All ages"):
            raise ValueError("NRS Table 1 columns changed")
        return [row(code, value, 2025) for name, code, area, sex, value in rows[1:]
                if area == "Council area" and sex == "Persons" and code in codes]
    finally:
        workbook.close()


def nrs_pcon(source, codes):
    expected = {code for code in codes if code.startswith("S140")}
    if len(expected) != 57:
        raise ValueError("Expected 57 current Scottish Hex constituencies")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["UKPC"]
        if sheet["A1"].value != (
                "Population estimates by UK Parliamentary Constituency (UKPC), sex and "
                "single year of age, mid-2011 to mid-2024 [note 7]"):
            raise ValueError("NRS UKPC geography/reference period changed; review the source")
        rows = sheet.iter_rows(min_row=4, max_col=5, values_only=True)
        if next(rows) != ("UKPC (2024) Name", "UKPC (2024) Code", "Sex", "Year", "All ages"):
            raise ValueError("NRS UKPC dimensions changed")
        selected = {}
        for name, code, sex, year, value in rows:
            if (name, code, sex, year, value) == (None,) * 5:
                continue
            if (code not in expected or sex not in {"Persons", "Males", "Females"}
                    or type(year) is not int or not 2011 <= year <= 2024):
                raise ValueError(f"Unexpected NRS UKPC dimensions: {code}, {sex}, {year}")
            if sex != "Persons" or year != 2024:
                continue
            if code in selected:
                raise ValueError(f"Duplicate NRS UKPC total: {code}")
            selected[code] = row(code, value, year)
        if selected.keys() != expected:
            raise ValueError(f"NRS UKPC coverage mismatch: missing {sorted(expected - selected.keys())}")
        return list(selected.values())
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_directory", type=Path)
    args = parser.parse_args()
    source = args.source_directory
    pcon = json.loads((ROOT / "data/PCON/uk-constituencies-2023.hexjson").read_text())["hexes"]
    lad = {f["properties"]["la_code"]: f["properties"]["la_name"] for f in json.loads(
        (ROOT / "data/LAD/uk_aq_la_hex_2025.geojson").read_text())["features"]}
    pcon_rows = nomis(source / "pcon.csv", pcon, 2024, "172")
    pcon_rows += nisra(source / "ni-pcon.xlsx", pcon,
                       "2. Parliamentary Constituencies (2024)", 2024)
    pcon_rows += nrs_pcon(source / "nrs-pcon.xlsx", pcon)
    lad_rows = nomis(source / "lad-current.csv", {c for c in lad if c[0] in "EW"}, 2025)
    lad_rows += nrs(source / "nrs.xlsx", lad)
    lad_rows += nisra(source / "ni-lad.xlsx", lad,
                      "2. Local Government Districts (LGD2014)", 2025)

    # Fail closed if coverage changes: review sources/boundaries before updating these gaps.
    expected_missing = {
        "PCON": set(),
        "LAD": {"E08000038", "E08000039"},
    }
    payloads = {}
    for geo, rows, codes in [("PCON", pcon_rows, pcon), ("LAD", lad_rows, lad)]:
        found = {r["geo_code"] for r in rows}
        if len(found) != len(rows) or found - codes.keys():
            raise ValueError(f"Duplicate/non-map {geo} codes")
        if codes.keys() - found != expected_missing[geo]:
            raise ValueError(f"Unexpected {geo} omissions: {sorted(codes.keys() - found)}")
        payloads[geo] = dict(geo_type=geo, count=len(rows), data=sorted(rows, key=lambda r: r["geo_code"]))
        print(f"{geo}: {len(rows)} rows; missing {len(codes) - len(found)}")

    destination = ROOT / "data/population"
    destination.mkdir(exist_ok=True)
    for geo, payload in payloads.items():
        # One compact row per line keeps the runtime files small and diffs reviewable.
        entries = ",\n".join("    " + json.dumps(r, separators=(",", ":")) for r in payload["data"])
        text = f'{{\n  "geo_type":"{geo}",\n  "count":{payload["count"]},\n  "data":[\n{entries}\n  ]\n}}\n'
        (destination / f"{geo.lower()}-latest.json").write_text(text)
    for name in ["pcon.csv", "lad-current.csv", "ni-pcon.xlsx", "ni-lad.xlsx", "nrs.xlsx", "nrs-pcon.xlsx"]:
        print(f"{name}: SHA-256 {hashlib.sha256((source / name).read_bytes()).hexdigest()}")


if __name__ == "__main__":
    main()
